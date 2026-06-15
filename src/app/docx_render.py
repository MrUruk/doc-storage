"""Best-effort HTML -> DOCX rendering.

The DOCX is a *derived* artifact for human delivery; content_html remains the
source of truth. Rendering runs as a separate step and is allowed to fail
without invalidating the saved version (the agent specification requires this).

This renderer covers the structures that appear in deploy orders: headings,
paragraphs, lists, and tables — including rowspan/colspan, which are preserved
by merging cells rather than dropping the structure.
"""

from __future__ import annotations

import asyncio
import io

from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}


async def render_docx(content_html: str) -> bytes:
    """Render HTML to DOCX bytes off the event loop."""
    return await asyncio.to_thread(_render_sync, content_html)


def _render_sync(content_html: str) -> bytes:
    try:
        soup = BeautifulSoup(content_html or "", "lxml")
    except Exception:
        soup = BeautifulSoup(content_html or "", "html.parser")

    document = Document()
    root = soup.body or soup
    for child in root.children:
        if isinstance(child, Tag):
            _render_block(document, child)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_block(document: Document, tag: Tag) -> None:
    name = tag.name.lower() if tag.name else ""
    if name in _HEADINGS:
        document.add_heading(tag.get_text(strip=True), level=_HEADINGS[name])
    elif name == "p":
        text = tag.get_text(" ", strip=True)
        if text:
            document.add_paragraph(text)
    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        for li in tag.find_all("li", recursive=False):
            document.add_paragraph(li.get_text(" ", strip=True), style=style)
    elif name == "table":
        _render_table(document, tag)
    elif name in ("div", "section", "article", "header", "footer", "main"):
        # Recurse into structural containers.
        for child in tag.children:
            if isinstance(child, Tag):
                _render_block(document, child)
    else:
        text = tag.get_text(" ", strip=True)
        if text:
            document.add_paragraph(text)


def _render_table(document: Document, table: Tag) -> None:
    rows = _collect_rows(table)
    if not rows:
        return
    n_cols = max((sum(_colspan(c) for c in row) for row in rows), default=0)
    n_rows = len(rows)
    if n_cols == 0:
        return

    docx_table = document.add_table(rows=n_rows, cols=n_cols)
    docx_table.style = "Table Grid"

    # occupied[r][c] marks grid cells already consumed by a span.
    occupied = [[False] * n_cols for _ in range(n_rows)]
    for r, row in enumerate(rows):
        c = 0
        for cell in row:
            while c < n_cols and occupied[r][c]:
                c += 1
            if c >= n_cols:
                break
            colspan = min(_colspan(cell), n_cols - c)
            rowspan = min(_rowspan(cell), n_rows - r)
            text = cell.get_text(" ", strip=True)

            top_left = docx_table.cell(r, c)
            top_left.text = text
            if cell.name and cell.name.lower() == "th":
                for paragraph in top_left.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True

            # Merge the spanned region into the top-left cell.
            bottom_right = docx_table.cell(
                r + rowspan - 1, c + colspan - 1
            )
            if bottom_right is not top_left:
                top_left.merge(bottom_right)
            for rr in range(r, r + rowspan):
                for cc in range(c, c + colspan):
                    occupied[rr][cc] = True
            c += colspan


def _collect_rows(table: Tag) -> list[list[Tag]]:
    rows: list[list[Tag]] = []
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"], recursive=False)
        if cells:
            rows.append(cells)
    return rows


def _colspan(cell: Tag) -> int:
    return _int_attr(cell, "colspan")


def _rowspan(cell: Tag) -> int:
    return _int_attr(cell, "rowspan")


def _int_attr(cell: Tag, name: str) -> int:
    try:
        return max(1, int(cell.get(name, "1")))
    except (TypeError, ValueError):
        return 1
